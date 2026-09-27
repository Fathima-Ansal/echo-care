import os
import subprocess
import sys

def install_and_import(package):
    try:
        import docx
    except ImportError:
        print(f"Installing {package}...")
        subprocess.check_call([sys.executable, "-m", "pip", "install", package])
        import docx
    return docx

docx = install_and_import('python-docx')
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_doc():
    doc = Document()
    
    # Title
    title = doc.add_heading('EchoCare: Complete API Workflow Guide', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Overview
    doc.add_heading('1. Overview of Major Libraries', level=1)
    
    p = doc.add_paragraph()
    p.add_run('Frontend Libraries:\n').bold = True
    p.add_run('• React & React-DOM: Core UI framework.\n')
    p.add_run('• React-Router-DOM: Client-side routing for navigating between pages.\n')
    p.add_run('• TailwindCSS: Utility-first CSS framework for styling.\n')
    p.add_run('• Lucide-React: Icon library for UI elements.\n')
    p.add_run('• Vite: Build tool and development server.\n')
    
    p = doc.add_paragraph()
    p.add_run('Backend Libraries:\n').bold = True
    p.add_run('• Flask & Flask-CORS: Web server framework and CORS handling.\n')
    p.add_run('• PyMongo: MongoDB database driver.\n')
    p.add_run('• Bcrypt: Secure password hashing.\n')
    p.add_run('• PyJWT: Authentication using JSON Web Tokens.\n')
    p.add_run('• Groq: LLM inference for generating empathetic conversational triage.\n')
    p.add_run('• SpeechRecognition & Pydub: Transcribing and processing user audio.\n')
    p.add_run('• gTTS (Google Text-to-Speech): Converting AI replies back to audio.\n')
    p.add_run('• Hume: Acoustic emotion recognition (analyzing voice prosody).\n')
    p.add_run('• Twilio: Initiating automated emergency SOS phone calls.\n')
    p.add_run('• python-dotenv: Environment variable management.\n')

    doc.add_page_break()

    # Endpoints
    doc.add_heading('2. Step-by-Step Endpoint Workflows', level=1)
    
    endpoints = [
        {
            "method": "POST",
            "path": "/api/register",
            "desc": "Registers a new caregiver user into the EchoCare system.",
            "libs": "React, Flask, Bcrypt, PyMongo",
            "steps": [
                ("Frontend (React)", "User enters their Email and Password in src/pages/Register.jsx."),
                ("Network (Fetch)", "Sends a POST request to /api/register with the user's details as a JSON payload."),
                ("Backend (Flask)", "The register() function is triggered to handle the request."),
                ("Security (Bcrypt)", "Hashes the plain-text password securely."),
                ("Database (PyMongo)", "Assigns the 'caregiver' role and inserts the new user document into the MongoDB users collection."),
                ("Response", "Backend returns a 201 Created status with a success message."),
                ("Navigation (React-Router)", "Redirects the user to the Login page to authenticate.")
            ]
        },
        {
            "method": "POST",
            "path": "/api/login",
            "desc": "Authenticates a user and provides a secure session token.",
            "libs": "React, React-Router-DOM, Flask, PyMongo, Bcrypt, PyJWT",
            "steps": [
                ("Frontend (React)", "User submits their credentials in src/pages/Login.jsx."),
                ("Network (Fetch)", "Sends a POST request to /api/login with the credentials."),
                ("Backend (Flask)", "The login() function is triggered."),
                ("Database (PyMongo)", "Searches MongoDB for a user matching the provided email."),
                ("Security (Bcrypt)", "Verifies the provided password against the stored hashed password."),
                ("Authentication (PyJWT)", "Generates a JWT Token containing the user_id and role."),
                ("Response", "Backend returns a 200 OK status containing the Token, User ID, Role, and Email."),
                ("Frontend State", "Saves the token and user details in the browser's localStorage."),
                ("Navigation (React-Router)", "Redirects the user to either the Caregiver Dashboard or Elderly Dashboard based on their role.")
            ]
        },
        {
            "method": "POST",
            "path": "/api/create-elderly",
            "desc": "Allows a caregiver to create an account for an elderly user linked to them.",
            "libs": "React, Flask, PyJWT, Bcrypt, PyMongo",
            "steps": [
                ("Frontend (React)", "Caregiver fills out the 'Add Elderly' form in src/pages/CaretakerDashboard.jsx."),
                ("Network (Fetch)", "Sends a POST request to /api/create-elderly. Includes the JWT Bearer Token in the headers for authorization."),
                ("Backend (Flask)", "The @token_required middleware intercepts the request."),
                ("Authentication (PyJWT)", "Validates the caregiver's token and confirms their role."),
                ("Backend (Flask)", "The create_elderly() function is triggered."),
                ("Security (Bcrypt)", "Hashes the new elderly user's password."),
                ("Database (PyMongo)", "Inserts the elderly user into MongoDB, linking them to the caregiver via caregiver_id."),
                ("Response", "Backend returns a 201 Created status."),
                ("Frontend State (React)", "Displays a success alert and updates the UI.")
            ]
        },
        {
            "method": "POST",
            "path": "/api/transcribe",
            "desc": "The core conversational AI loop handling audio input, transcription, LLM triage, and TTS response.",
            "libs": "React, Flask, PyJWT, SpeechRecognition, Groq, gTTS, Hume, PyMongo",
            "steps": [
                ("Frontend (React)", "Elderly user records a voice message in src/pages/ElderlyDashboard.jsx using MediaRecorder."),
                ("Network (Fetch)", "Sends audio as multipart/form-data via POST to /api/transcribe (with Bearer Token)."),
                ("Backend (Flask)", "Validates token and triggers the transcribe() function."),
                ("AI (SpeechRecognition)", "Converts the received audio file into text."),
                ("Database (PyMongo)", "Fetches the recent conversation history to provide context and prevent the AI from repeating questions."),
                ("AI (Groq Llama 3)", "Generates an empathetic, human-like triage response in JSON format (including reply text and base sentiment)."),
                ("AI (gTTS)", "Converts the generated AI reply text back into an audio file (Base64 string)."),
                ("AI (Hume)", "Analyzes the prosody (tone/emotion) of the original voice. Overrides the sentiment to 'Negative' if strong distress is detected."),
                ("Database (PyMongo)", "Saves the entire interaction log (user text, AI reply, sentiment, audio) into the health_logs collection."),
                ("Response", "Returns the transcribed text, AI reply, finalized sentiment, and Base64 audio."),
                ("Frontend (React)", "Displays the conversation text visually and automatically plays back the AI audio response.")
            ]
        },
        {
            "method": "GET",
            "path": "/api/logs",
            "desc": "Fetches conversation history logs for the dashboard.",
            "libs": "React, Flask, PyJWT, PyMongo",
            "steps": [
                ("Frontend (React)", "Dashboard component mounts and triggers a GET request to /api/logs (with Bearer Token)."),
                ("Backend (Flask)", "Validates the token and triggers get_logs()."),
                ("Logic (Flask)", "Checks the user role. If 'elderly', it prepares to fetch only their logs. If 'caregiver', it prepares to fetch logs for all elderly users linked to their caregiver_id."),
                ("Database (PyMongo)", "Retrieves the relevant logs and sorts them by timestamp."),
                ("Response", "Returns a JSON array of the formatted logs."),
                ("Frontend (React)", "Renders the logs into tables or lists for review.")
            ]
        },
        {
            "method": "POST",
            "path": "/api/caretaker/profile",
            "desc": "Updates the caregiver's emergency contact phone number.",
            "libs": "React, Flask, PyJWT, PyMongo",
            "steps": [
                ("Frontend (React)", "Caregiver submits their phone number in the settings tab of CaretakerDashboard.jsx."),
                ("Network (Fetch)", "Sends a POST request to /api/caretaker/profile (with Bearer Token)."),
                ("Backend (Flask)", "Validates the token and triggers update_caretaker_profile()."),
                ("Database (PyMongo)", "Updates the caregiver's user document with the new phone_number field."),
                ("Response", "Returns a 200 OK success message."),
                ("Frontend (React)", "Shows a confirmation alert to the user.")
            ]
        },
        {
            "method": "POST",
            "path": "/api/sos",
            "desc": "Triggers an emergency automated phone call to the caregiver.",
            "libs": "React, Flask, PyJWT, PyMongo, Twilio",
            "steps": [
                ("Frontend (React)", "Elderly user clicks the prominent Red SOS button in src/pages/ElderlyDashboard.jsx."),
                ("Network (Fetch)", "Sends a POST request to /api/sos (with Bearer Token)."),
                ("Backend (Flask)", "Validates the token and triggers trigger_sos()."),
                ("Database (PyMongo)", "Fetches the elderly user's record to find their associated caregiver_id, then fetches the caregiver's phone_number."),
                ("API (Twilio)", "Uses the Twilio API to initiate a voice call to the caregiver's phone, playing an automated emergency TTS message."),
                ("Response", "Returns a success confirmation to the frontend."),
                ("Frontend (React)", "Displays visual feedback that the SOS call was successfully placed.")
            ]
        }
    ]

    for ep in endpoints:
        # Endpoint Title
        head = doc.add_heading(f"{ep['method']} {ep['path']}", level=2)
        
        # Description
        p_desc = doc.add_paragraph()
        p_desc.add_run("Description: ").bold = True
        p_desc.add_run(ep['desc'])
        
        # Libraries
        p_libs = doc.add_paragraph()
        p_libs.add_run("Major Libraries: ").bold = True
        p_libs.add_run(ep['libs'])
        
        # Steps
        p_steps = doc.add_paragraph()
        p_steps.add_run("Step-by-Step Flow:").bold = True
        
        for i, (actor, action) in enumerate(ep['steps']):
            p_step = doc.add_paragraph(style='List Number')
            # Make the actor/component bold
            run = p_step.add_run(f"{actor}: ")
            run.bold = True
            # Add the action
            p_step.add_run(action)
            
        doc.add_paragraph() # Add some spacing between endpoints

    # Save
    doc.save('EchoCare_Detailed_API_Workflow_v2.docx')
    print("Successfully created EchoCare_Detailed_API_Workflow_v2.docx")

if __name__ == '__main__':
    create_doc()
